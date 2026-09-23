import os
from datetime import date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_schedule():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Training Schedule"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Color Palette & Styles (Premium Corporate Theme)
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    stripe_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Slate 50
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    bold_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    regular_font = Font(name="Calibri", size=10, color="1E293B")
    
    thin_border_side = Side(style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Column Headers expected by the platform's ExcelIngestionService
    headers = [
        "Batch ID",
        "Date of Training",
        "Day",
        "Start Time",
        "End Time",
        "No of Hours",
        "Topic",
        "Faculty Name",
        "Mode of Delivery",
        "Location City",
        "Venue"
    ]

    ws.append(headers)

    # Style Header Row
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = border_cell

    ws.row_dimensions[1].height = 28

    # Schedule data for January 2027 (Next Year January)
    # 15 delivery days (3-week curriculum, Monday to Friday)
    schedule_data = [
        # Week 1: Foundation & Core Architecture
        ("BATCH_123", date(2027, 1, 4), "Monday", "09:30", "17:30", 8, "Module 1: Orientation, Cloud Infrastructure & Distributed Architecture Overview", "Dr. Rajesh Kumar", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 5), "Tuesday", "09:30", "17:30", 8, "Module 2: PySpark Core - Resilient Distributed Datasets (RDDs) & Memory Management", "Dr. Rajesh Kumar", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 6), "Wednesday", "09:30", "17:30", 8, "Module 3: DataFrames & Spark SQL - Catalyst Query Optimizer & Tungsten Engine", "Dr. Rajesh Kumar", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 7), "Thursday", "09:30", "17:30", 8, "Module 4: Advanced Transformations, Window Functions & Broadcast Joins", "Dr. Rajesh Kumar", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 8), "Friday", "09:30", "17:30", 8, "Module 5: Performance Tuning, Partitioning, Caching & Spill Diagnostics", "Dr. Rajesh Kumar", "Online", "Bengaluru", "MS Teams Virtual Room A"),

        # Week 2: Lakehouse & Data Engineering Pipelines
        ("BATCH_123", date(2027, 1, 11), "Monday", "09:30", "17:30", 8, "Module 6: Databricks Workspace, Compute Clusters & Delta Lake Architecture", "Amit Sharma", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 12), "Tuesday", "09:30", "17:30", 8, "Module 7: ACID Transactions, Time Travel, Schema Evolution & Compaction in Delta", "Amit Sharma", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 13), "Wednesday", "09:30", "17:30", 8, "Module 8: Real-Time Structured Streaming & Event Ingestion with Kafka Integration", "Amit Sharma", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 14), "Thursday", "09:30", "17:30", 8, "Module 9: Delta Live Tables (DLT) - Declarative ETL Pipelines & Data Quality Expectations", "Amit Sharma", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 15), "Friday", "09:30", "17:30", 8, "Module 10: Databricks Asset Bundles (DABs), CI/CD & Production Multi-Task Workflows", "Amit Sharma", "Online", "Bengaluru", "MS Teams Virtual Room A"),

        # Week 3: Governance, Security, ML & Capstone Delivery
        ("BATCH_123", date(2027, 1, 18), "Monday", "09:30", "17:30", 8, "Module 11: Unity Catalog - Fine-Grained Access Control, Data Lineage & Governance", "Dr. Rajesh Kumar", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 19), "Tuesday", "09:30", "17:30", 8, "Module 12: Machine Learning on Databricks, Feature Store & MLflow Experiment Tracking", "Amit Sharma", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 20), "Wednesday", "09:30", "17:30", 8, "Module 13: End-to-End Enterprise Lakehouse Capstone Architecture & Data Modeling", "Dr. Rajesh Kumar", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 21), "Thursday", "09:30", "17:30", 8, "Module 14: Capstone Pipeline Implementation, Conflict Resolution & Peer Review", "Dr. Rajesh Kumar", "Online", "Bengaluru", "MS Teams Virtual Room A"),
        ("BATCH_123", date(2027, 1, 22), "Friday", "09:30", "17:30", 8, "Module 15: Final Capstone Demonstration, Benchmark Evaluation & Quality Closure (Gate 1)", "Dr. Rajesh Kumar", "Online", "Bengaluru", "MS Teams Virtual Room A"),
    ]

    for row_idx, row_data in enumerate(schedule_data, start=2):
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 22
        is_stripe = (row_idx % 2 == 1)
        row_fill = stripe_fill if is_stripe else white_fill

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = row_fill
            cell.font = regular_font
            cell.border = border_cell

            # Alignments & formatting
            if col_idx in (1, 2, 3, 4, 5, 9):
                cell.alignment = align_center
            elif col_idx == 6:
                cell.alignment = align_right
            else:
                cell.alignment = align_left

            if col_idx == 1:
                cell.font = bold_font
            
            # Format Date column specifically as DD-MM-YYYY
            if col_idx == 2:
                cell.number_format = "DD-MM-YYYY"

    # Specific fine-tuned column widths
    ws.column_dimensions["A"].width = 16  # Batch ID
    ws.column_dimensions["B"].width = 18  # Date of Training
    ws.column_dimensions["C"].width = 14  # Day
    ws.column_dimensions["D"].width = 14  # Start Time
    ws.column_dimensions["E"].width = 14  # End Time
    ws.column_dimensions["F"].width = 14  # No of Hours
    ws.column_dimensions["G"].width = 62  # Topic
    ws.column_dimensions["H"].width = 24  # Faculty Name
    ws.column_dimensions["I"].width = 18  # Mode of Delivery
    ws.column_dimensions["J"].width = 18  # Location City
    ws.column_dimensions["K"].width = 30  # Venue

    # Save to primary locations
    output_paths = [
        "/mnt/d/projects and files/ops/batch_schedule_january_2027.xlsx",
        "/mnt/d/projects and files/ops/frontend/public/batch_schedule_january_2027.xlsx"
    ]

    for path in output_paths:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        wb.save(path)
        print(f"Schedule saved to: {path}")

if __name__ == "__main__":
    generate_schedule()
